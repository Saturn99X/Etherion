from __future__ import annotations

import hashlib
import logging
import mimetypes
import os
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Any, List, Optional

logger = logging.getLogger(__name__)

try:
    import tiktoken
except Exception:  # pragma: no cover
    tiktoken = None


@dataclass
class ExtractedImage:
    image_bytes: bytes
    mime_type: str
    chapter_heading: Optional[str] = None
    page_number: Optional[int] = None
    description: Optional[str] = None

    @property
    def content_hash(self) -> str:
        return hashlib.sha256(self.image_bytes).hexdigest()[:16]


@dataclass
class ChapterEssence:
    heading: str
    level: int
    start_line: int
    essence_text: str
    full_content: str
    images: List[ExtractedImage] = field(default_factory=list)


@dataclass
class DocumentParseResult:
    markdown: str
    chapters: List[ChapterEssence]
    images: List[ExtractedImage]
    total_chars: int
    estimated_tokens: int
    source_filename: str
    mime_type: str
    metadata: dict = field(default_factory=dict)


class PyMuPDFParserService:
    MAX_ESSENCE_TOKENS = 128

    @staticmethod
    @lru_cache(maxsize=1)
    def _encoder():
        if tiktoken is None:
            raise RuntimeError("tiktoken is required for token counting")
        return tiktoken.get_encoding("cl100k_base")

    def _count_tokens(self, text: str) -> int:
        if not text:
            return 0
        enc = self._encoder()
        return len(enc.encode(text))

    def _truncate_to_tokens(self, text: str, max_tokens: int) -> str:
        if not text:
            return ""
        enc = self._encoder()
        prefix_chars = int(os.getenv("INGEST_TOKEN_TRUNCATE_PREFIX_CHARS", "8192") or "8192")
        if len(text) > prefix_chars:
            toks = enc.encode(text[:prefix_chars])
            if len(toks) >= max_tokens:
                return enc.decode(toks[:max_tokens])
        toks = enc.encode(text)
        if len(toks) <= max_tokens:
            return text
        return enc.decode(toks[:max_tokens])

    def parse_bytes(self, file_bytes: bytes, filename: str, mime_type: Optional[str] = None) -> DocumentParseResult:
        suffix = os.path.splitext(filename)[1].lower() if "." in filename else ""
        effective_mime = (mime_type or self._guess_mime_type(filename)).split(";", 1)[0].strip().lower()

        if suffix in {".docx", ".pptx", ".xlsx"} or effective_mime in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }:
            return self._parse_office_bytes(file_bytes=file_bytes, filename=filename, mime_type=effective_mime, suffix=suffix)

        if effective_mime.startswith("text/") or effective_mime == "application/json" or suffix in {
            ".txt",
            ".md",
            ".csv",
            ".log",
            ".json",
            ".yaml",
            ".yml",
            ".html",
            ".htm",
        }:
            text = file_bytes.decode("utf-8", errors="replace")
            total_chars = len(text)
            est_tokens = self._count_tokens(text)
            heading = filename or "Document"
            essence_body = self._truncate_to_tokens(text, self.MAX_ESSENCE_TOKENS).strip()
            essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading
            ch = ChapterEssence(
                heading=heading,
                level=1,
                start_line=0,
                essence_text=essence_text,
                full_content=text,
                images=[],
            )
            return DocumentParseResult(
                markdown=text,
                chapters=[ch],
                images=[],
                total_chars=total_chars,
                estimated_tokens=est_tokens,
                source_filename=filename,
                mime_type=effective_mime,
                metadata={"parser": "pymupdf_text", "chapter_count": 1, "image_count": 0},
            )

        # Prefer explicit PDF parsing path.
        if effective_mime == "application/pdf" or suffix == ".pdf":
            return self._parse_document(file_bytes=file_bytes, filename=filename, mime_type=effective_mime, filetype_hint="pdf")

        if effective_mime.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tiff", ".bmp"}:
            img = ExtractedImage(
                image_bytes=file_bytes,
                mime_type=effective_mime if effective_mime.startswith("image/") else self._guess_mime_type(filename),
                chapter_heading=filename,
                page_number=None,
                description=None,
            )
            ch = ChapterEssence(
                heading=filename,
                level=1,
                start_line=0,
                essence_text=filename,
                full_content="",
                images=[img],
            )
            return DocumentParseResult(
                markdown="",
                chapters=[ch],
                images=[img],
                total_chars=0,
                estimated_tokens=0,
                source_filename=filename,
                mime_type=effective_mime,
                metadata={"parser": "pymupdf", "chapter_count": 1, "image_count": 1},
            )

        # For other PyMuPDF-supported formats (xps, epub, mobi, fb2, cbz, svg, etc.)
        filetype_hint = suffix[1:] if suffix.startswith(".") and len(suffix) > 1 else None
        return self._parse_document(file_bytes=file_bytes, filename=filename, mime_type=effective_mime, filetype_hint=filetype_hint)

    def _parse_office_bytes(
        self,
        *,
        file_bytes: bytes,
        filename: str,
        mime_type: str,
        suffix: str,
    ) -> DocumentParseResult:
        if suffix == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self._parse_docx(file_bytes=file_bytes, filename=filename, mime_type=mime_type)
        if suffix == ".pptx" or mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return self._parse_pptx(file_bytes=file_bytes, filename=filename, mime_type=mime_type)
        if suffix == ".xlsx" or mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return self._parse_xlsx(file_bytes=file_bytes, filename=filename, mime_type=mime_type)

        # Fallback to generic text parsing if unknown office type
        text = file_bytes.decode("utf-8", errors="replace")
        total_chars = len(text)
        est_tokens = self._count_tokens(text)
        heading = filename or "Document"
        essence_body = self._truncate_to_tokens(text, self.MAX_ESSENCE_TOKENS).strip()
        essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading
        ch = ChapterEssence(heading=heading, level=1, start_line=0, essence_text=essence_text, full_content=text, images=[])
        return DocumentParseResult(
            markdown=text,
            chapters=[ch],
            images=[],
            total_chars=total_chars,
            estimated_tokens=est_tokens,
            source_filename=filename,
            mime_type=mime_type,
            metadata={"parser": "office_fallback", "chapter_count": 1, "image_count": 0},
        )

    def _parse_docx(self, *, file_bytes: bytes, filename: str, mime_type: str) -> DocumentParseResult:
        import io

        try:
            from docx import Document  # python-docx
        except ModuleNotFoundError as e:
            raise RuntimeError("python-docx is required to parse .docx") from e

        doc = Document(io.BytesIO(file_bytes))

        lines: List[str] = []
        for p in getattr(doc, "paragraphs", []) or []:
            t = (p.text or "").strip()
            if t:
                lines.append(t)

        # Include basic table text
        for table in getattr(doc, "tables", []) or []:
            for row in getattr(table, "rows", []) or []:
                cells = []
                for cell in getattr(row, "cells", []) or []:
                    v = (cell.text or "").strip()
                    cells.append(v)
                if any(cells):
                    lines.append("\t".join(cells))

        text = "\n".join(lines).strip()
        total_chars = len(text)
        est_tokens = self._count_tokens(text)
        heading = filename or "Document"
        essence_body = self._truncate_to_tokens(text, self.MAX_ESSENCE_TOKENS).strip()
        essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading

        ch = ChapterEssence(heading=heading, level=1, start_line=0, essence_text=essence_text, full_content=text, images=[])
        return DocumentParseResult(
            markdown=text,
            chapters=[ch],
            images=[],
            total_chars=total_chars,
            estimated_tokens=est_tokens,
            source_filename=filename,
            mime_type=mime_type,
            metadata={"parser": "python_docx", "chapter_count": 1, "image_count": 0},
        )

    def _parse_pptx(self, *, file_bytes: bytes, filename: str, mime_type: str) -> DocumentParseResult:
        import io

        try:
            from pptx import Presentation  # python-pptx
        except ModuleNotFoundError as e:
            raise RuntimeError("python-pptx is required to parse .pptx") from e

        pres = Presentation(io.BytesIO(file_bytes))

        chapters: List[ChapterEssence] = []
        markdown_parts: List[str] = []
        total_chars = 0
        total_tokens = 0

        for i, slide in enumerate(getattr(pres, "slides", []) or []):
            slide_no = i + 1
            heading = f"Slide {slide_no}"
            lines: List[str] = []
            for shape in getattr(slide, "shapes", []) or []:
                try:
                    if getattr(shape, "has_text_frame", False):
                        t = (shape.text or "").strip()
                        if t:
                            lines.append(t)
                except Exception:
                    continue
            slide_text = "\n".join(lines).strip()
            total_chars += len(slide_text)
            total_tokens += self._count_tokens(slide_text)

            essence_body = self._truncate_to_tokens(slide_text, self.MAX_ESSENCE_TOKENS).strip()
            essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading
            chapters.append(
                ChapterEssence(
                    heading=heading,
                    level=1,
                    start_line=i,
                    essence_text=essence_text,
                    full_content=slide_text,
                    images=[],
                )
            )
            if slide_text:
                markdown_parts.append(f"## {heading}\n\n{slide_text}")

        markdown = "\n\n".join(markdown_parts)
        return DocumentParseResult(
            markdown=markdown,
            chapters=chapters,
            images=[],
            total_chars=total_chars,
            estimated_tokens=total_tokens,
            source_filename=filename,
            mime_type=mime_type,
            metadata={"parser": "python_pptx", "chapter_count": len(chapters), "image_count": 0},
        )

    def _parse_xlsx(self, *, file_bytes: bytes, filename: str, mime_type: str) -> DocumentParseResult:
        import io

        try:
            import openpyxl
        except ModuleNotFoundError as e:
            raise RuntimeError("openpyxl is required to parse .xlsx") from e

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

        chapters: List[ChapterEssence] = []
        markdown_parts: List[str] = []
        total_chars = 0
        total_tokens = 0

        for sheet_idx, sheet in enumerate(wb.worksheets):
            heading = f"Sheet: {sheet.title}" if getattr(sheet, "title", None) else f"Sheet {sheet_idx+1}"
            rows: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    rows.append("\t".join(vals))
            sheet_text = "\n".join(rows).strip()
            total_chars += len(sheet_text)
            total_tokens += self._count_tokens(sheet_text)

            essence_body = self._truncate_to_tokens(sheet_text, self.MAX_ESSENCE_TOKENS).strip()
            essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading
            chapters.append(
                ChapterEssence(
                    heading=heading,
                    level=1,
                    start_line=sheet_idx,
                    essence_text=essence_text,
                    full_content=sheet_text,
                    images=[],
                )
            )
            if sheet_text:
                markdown_parts.append(f"## {heading}\n\n{sheet_text}")

        markdown = "\n\n".join(markdown_parts)
        return DocumentParseResult(
            markdown=markdown,
            chapters=chapters,
            images=[],
            total_chars=total_chars,
            estimated_tokens=total_tokens,
            source_filename=filename,
            mime_type=mime_type,
            metadata={"parser": "openpyxl", "chapter_count": len(chapters), "image_count": 0},
        )

    def _parse_document(
        self,
        *,
        file_bytes: bytes,
        filename: str,
        mime_type: str,
        filetype_hint: Optional[str],
    ) -> DocumentParseResult:
        try:
            import fitz  # PyMuPDF
        except ModuleNotFoundError as e:
            raise RuntimeError(
                "PyMuPDF is required for deterministic document parsing. "
                "Install it in your active venv: `venv/bin/python -m pip install PyMuPDF>=1.24.0`."
            ) from e

        # PyMuPDF can open many formats; providing a filetype hint improves determinism.
        # If the hint is wrong, fall back to auto-detection.
        try:
            if filetype_hint:
                doc = fitz.open(stream=file_bytes, filetype=filetype_hint)
            else:
                doc = fitz.open(stream=file_bytes)
        except Exception:
            doc = fitz.open(stream=file_bytes)

        chapters: List[ChapterEssence] = []
        images: List[ExtractedImage] = []
        seen_hashes: set[str] = set()

        markdown_pages: List[str] = []
        total_chars = 0

        try:
            max_pages_env = int(os.getenv("INGEST_MAX_DOC_CHUNKS", "0") or "0")
        except Exception:
            max_pages_env = 0
        page_limit = doc.page_count
        if max_pages_env > 0:
            page_limit = min(page_limit, max_pages_env)

        fast_pdf_essence = os.getenv("INGEST_FAST_PDF_ESSENCE", "").lower() in {"1", "true", "yes"}

        extract_images = os.getenv("INGEST_PDF_EXTRACT_IMAGES", "0").lower() in {"1", "true", "yes"}
        if not extract_images and doc.page_count <= 64:
            extract_images = False

        for page_idx in range(page_limit):
            page = doc.load_page(page_idx)
            page_number = page_idx + 1
            heading = f"Page {page_number}"

            text = (page.get_text("text") or "").strip()
            total_chars += len(text)

            if fast_pdf_essence:
                essence_body = (text[:1024] or "").strip()
            else:
                essence_body = self._truncate_to_tokens(text, self.MAX_ESSENCE_TOKENS).strip()
            essence_text = f"{heading}\n\n{essence_body}".strip() if essence_body else heading

            page_images: List[ExtractedImage] = []
            try:
                if extract_images:
                    for img in page.get_images(full=True) or []:
                        xref = img[0]
                        base = doc.extract_image(xref)
                        if not base:
                            continue
                        img_bytes = base.get("image")
                        if not img_bytes:
                            continue
                        ext = (base.get("ext") or "png").lower()
                        img_mime = f"image/{ext}" if ext else "image/png"

                        h = hashlib.sha256(img_bytes).hexdigest()[:16]
                        if h in seen_hashes:
                            continue
                        seen_hashes.add(h)

                        ex = ExtractedImage(
                            image_bytes=img_bytes,
                            mime_type=img_mime,
                            chapter_heading=heading,
                            page_number=page_number,
                            description=None,
                        )
                        images.append(ex)
                        page_images.append(ex)
            except Exception as e:
                logger.warning(f"PyMuPDF image extraction failed for {filename} page {page_number}: {e}")

            chapters.append(
                ChapterEssence(
                    heading=heading,
                    level=1,
                    start_line=page_idx,
                    essence_text=essence_text,
                    full_content=text,
                    images=page_images,
                )
            )

            if text:
                markdown_pages.append(f"## {heading}\n\n{text}")

        markdown = "\n\n".join(markdown_pages)
        try:
            chars_per_token = int(os.getenv("INGEST_ESTIMATED_CHARS_PER_TOKEN", "4") or "4")
        except Exception:
            chars_per_token = 4
        estimated_tokens = int(total_chars / max(1, chars_per_token))

        return DocumentParseResult(
            markdown=markdown,
            chapters=chapters,
            images=images,
            total_chars=total_chars,
            estimated_tokens=estimated_tokens,
            source_filename=filename,
            mime_type=mime_type,
            metadata={"parser": "pymupdf", "chapter_count": len(chapters), "image_count": len(images)},
        )

    def _guess_mime_type(self, filename: str) -> str:
        guess, _ = mimetypes.guess_type(filename)
        return guess or "application/octet-stream"


def should_split_document(estimated_tokens: int, max_tokens: int = 300_000) -> bool:
    return estimated_tokens >= max_tokens


def split_chapters_into_parts(
    chapters: List[ChapterEssence],
    filename: str,
    max_tokens_per_part: int = 300_000,
    chars_per_token: int = 4,
) -> List[dict]:
    parts: List[dict] = []
    current_part: dict = {"chapters": [], "char_count": 0, "token_count": 0}

    if tiktoken is None:
        raise RuntimeError("tiktoken is required for splitting")
    enc = tiktoken.get_encoding("cl100k_base")

    for chapter in chapters:
        chapter_chars = len(chapter.full_content)
        chapter_tokens = len(enc.encode(chapter.full_content or ""))

        if current_part["token_count"] + chapter_tokens > max_tokens_per_part and current_part["chapters"]:
            parts.append(current_part)
            current_part = {"chapters": [], "char_count": 0, "token_count": 0}

        current_part["chapters"].append(chapter)
        current_part["char_count"] += chapter_chars
        current_part["token_count"] += chapter_tokens

    if current_part["chapters"]:
        parts.append(current_part)

    total = len(parts)
    base_name = filename.rsplit(".", 1)[0] if "." in filename else filename
    ext = filename.rsplit(".", 1)[1] if "." in filename else ""

    for i, part in enumerate(parts):
        part_number = i + 1
        part["part_number"] = part_number
        part["total_parts"] = total
        part["part_name"] = f"{base_name}_{part_number:03d}.{ext}" if ext else f"{base_name}_{part_number:03d}"

    return parts
